"""
Frontend/app.py

Streamlit UI for LegalEase. Collects document details, sends them to
the FastAPI backend's /generate endpoint, previews the AI-generated
document, allows inline editing, and offers .txt/.docx/.pdf downloads.

Run with:
    streamlit run app.py   (run from inside Frontend/)
"""

import os
import re
from io import BytesIO

import requests
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF

API_BASE_URL = os.getenv("API_BASE_URL", "http://localhost:8000")
FOOTER_TEXT = "LegalEase Inc. | contact@legalease.com | All Rights Reserved."
LOGO_PATH = os.path.join(os.path.dirname(__file__), "Logo.png")  # optional, add your own


# --------------------------------------------------------------------------
# Formatting helpers
# --------------------------------------------------------------------------
def sanitize_text(text: str) -> str:
    if not text:
        return ""
    replacements = {
        "\u2018": "'", "\u2019": "'",
        "\u201c": '"', "\u201d": '"',
        "\u2013": "-", "\u2014": "-",
        "\u2026": "...", "\u2022": "-",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"[^\x00-\x7F]+", "", text)
    return text.strip()


def _split_terms(terms: str):
    return [t.strip() for t in terms.split(";") if t.strip()]


def _is_heading(raw_line: str, clean_line: str) -> bool:
    if raw_line.startswith("#"):
        return True
    words = clean_line.split()
    return bool(clean_line) and clean_line.isupper() and len(words) <= 8


def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    text = sanitize_text(text)
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    if os.path.exists(LOGO_PATH):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(LOGO_PATH, width=Inches(1.3))

    title = document.add_heading(doc_type or "Legal Document", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            document.add_paragraph("")
            continue
        clean_line = line.lstrip("#").strip()
        if _is_heading(line, clean_line):
            document.add_heading(clean_line, level=2)
        else:
            document.add_paragraph(clean_line)

    terms_list = _split_terms(terms)
    if terms_list:
        document.add_heading("Summary of Terms", level=2)
        table = document.add_table(rows=1, cols=2)
        try:
            table.style = "Light Grid Accent 1"
        except KeyError:
            pass
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text = "No.", "Term"
        for idx, term in enumerate(terms_list, start=1):
            row = table.add_row().cells
            row[0].text, row[1].text = str(idx), term

    document.add_paragraph("")
    footer = document.add_paragraph(FOOTER_TEXT)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.italic = True
        run.font.size = Pt(8)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class _BrandedPDF(FPDF):
    def __init__(self, doc_type: str):
        super().__init__()
        self.doc_type = doc_type or "Legal Document"
        self.set_auto_page_break(auto=True, margin=20)

    def header(self):
        if os.path.exists(LOGO_PATH):
            self.image(LOGO_PATH, x=(self.w - 22) / 2, y=8, w=22)
            self.ln(22)
        else:
            self.ln(6)
        self.set_font("Times", "B", 15)
        self.cell(0, 10, self.doc_type, align="C", ln=True)
        self.ln(3)

    def footer(self):
        self.set_y(-15)
        self.set_font("Times", "I", 8)
        self.cell(0, 10, FOOTER_TEXT, align="C")


def format_pdf(text: str, doc_type: str) -> bytes:
    text = sanitize_text(text)
    pdf = _BrandedPDF(doc_type)
    pdf.add_page()
    pdf.set_font("Times", size=11)
    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            pdf.ln(4)
            continue
        clean_line = line.lstrip("#").strip()
        if _is_heading(line, clean_line):
            pdf.set_font("Times", "B", 12)
            pdf.multi_cell(0, 8, clean_line)
            pdf.set_font("Times", size=11)
        elif line.startswith(("-", "*")):
            pdf.multi_cell(0, 7, f"    - {line.lstrip('-* ').strip()}")
        else:
            pdf.multi_cell(0, 7, clean_line)
    return bytes(pdf.output(dest="S"))


def format_html_preview(text: str) -> str:
    text = sanitize_text(text)
    parts, in_list = [], False
    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            if in_list:
                parts.append("</ul>")
                in_list = False
            parts.append("<br>")
            continue
        clean_line = line.lstrip("#").strip()
        if line.startswith(("-", "*")):
            if not in_list:
                parts.append("<ul style='margin:4px 0;'>")
                in_list = True
            parts.append(f"<li>{line.lstrip('-* ').strip()}</li>")
            continue
        if in_list:
            parts.append("</ul>")
            in_list = False
        if _is_heading(line, clean_line):
            parts.append(f"<h4 style='color:#e8c874;margin-top:16px;'>{clean_line}</h4>")
        else:
            parts.append(f"<p style='margin:4px 0;line-height:1.5;'>{clean_line}</p>")
    if in_list:
        parts.append("</ul>")
    return "".join(parts)


# --------------------------------------------------------------------------
# Streamlit UI
# --------------------------------------------------------------------------
st.set_page_config(page_title="LegalEase", layout="centered")

col1, col2, col3 = st.columns([1, 2, 1])
with col2:
    if os.path.exists(LOGO_PATH):
        st.image(LOGO_PATH, use_container_width=True)
    else:
        st.markdown("<h1 style='text-align:center;'>⚖️ LegalEase</h1>", unsafe_allow_html=True)

st.markdown("<h2 style='text-align:center;'>AI Legal Document Generator</h2>", unsafe_allow_html=True)

if "generated_text" not in st.session_state:
    st.session_state.generated_text = ""
if "show_edit" not in st.session_state:
    st.session_state.show_edit = False

document_type = st.text_input("Document Type (Ex: Agreement, Contract, NDA)")
parties = st.text_area("Parties Involved")
terms = st.text_area("Terms & Conditions (Use semicolons for bullet points)")
dates = st.text_input("Effective Date")

if st.button("Generate Document"):
    if not (document_type and parties and terms and dates):
        st.warning("Please fill in all fields before generating the document.")
    else:
        with st.spinner("Generating your legal document..."):
            try:
                response = requests.post(
                    f"{API_BASE_URL}/generate",
                    json={
                        "document_type": document_type,
                        "parties": parties,
                        "terms": terms,
                        "dates": dates,
                    },
                    timeout=60,
                )
                if response.status_code == 200:
                    st.session_state.generated_text = sanitize_text(response.json()["document"])
                    st.session_state.show_edit = False
                    st.success("Document Generated Successfully!")
                else:
                    try:
                        detail = response.json().get("detail", response.text)
                    except ValueError:
                        detail = response.text
                    st.error(f"Generation failed: {detail}")
            except requests.exceptions.ConnectionError:
                st.error(f"Could not connect to the backend API at {API_BASE_URL}. Is it running?")
            except Exception as exc:
                st.error(f"Unexpected error: {exc}")

if st.session_state.generated_text:
    st.markdown(
        f"""<div style="background-color:#0e1117; border:1px solid #333; border-radius:8px;
                    padding:20px; max-height:400px; overflow-y:auto; color:#e6e6e6;">
            {format_html_preview(st.session_state.generated_text)}
        </div>""",
        unsafe_allow_html=True,
    )

    if st.button("✏️ Click to Edit Document"):
        st.session_state.show_edit = not st.session_state.show_edit

    if st.session_state.show_edit:
        st.session_state.generated_text = st.text_area(
            "Edit Document Below:", st.session_state.generated_text, height=300
        )

    safe_name = (document_type or "document").strip().replace(" ", "_").lower() or "document"

    c1, c2, c3 = st.columns(3)
    with c1:
        st.download_button("📄 Download as .TXT", data=st.session_state.generated_text,
                            file_name=f"{safe_name}.txt", mime="text/plain", use_container_width=True)
    with c2:
        st.download_button("📝 Download as .DOCX",
                            data=format_docx(st.session_state.generated_text, document_type, terms),
                            file_name=f"{safe_name}.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            use_container_width=True)
    with c3:
        st.download_button("🧾 Download as .PDF",
                            data=format_pdf(st.session_state.generated_text, document_type),
                            file_name=f"{safe_name}.pdf", mime="application/pdf", use_container_width=True)
else:
    st.info("👉 Fill in the fields above and click 'Generate Document' to start.")